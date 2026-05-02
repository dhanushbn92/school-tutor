from io import BytesIO

from docx import Document
from docx.shared import Pt, RGBColor

from app.llm.schemas.lesson_plan import LessonPlanOutput


def render_lesson_plan_docx(plan: LessonPlanOutput) -> bytes:
    doc = Document()

    _add_heading(doc, plan.title, level=0)
    meta = doc.add_paragraph()
    meta_run = meta.add_run(
        f"Class {plan.class_level}  ·  {plan.subject}  ·  {plan.chapter_title}  ·  {plan.duration_minutes} min"
    )
    meta_run.italic = True
    meta_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    _bullet_section(doc, "Learning objectives", plan.objectives)
    _bullet_section(doc, "Prerequisites", plan.prerequisites)
    _bullet_section(doc, "Materials required", plan.materials)
    _bullet_section(doc, "Key vocabulary", plan.key_vocabulary)

    _add_heading(doc, "Lesson flow", level=1)
    for index, activity in enumerate(plan.activities, start=1):
        heading = doc.add_paragraph()
        h_run = heading.add_run(
            f"{index}. {activity.phase}  ({activity.duration_minutes} min)"
        )
        h_run.bold = True
        h_run.font.size = Pt(12)
        doc.add_paragraph(activity.description)
        _sub_bullets(doc, "Teacher:", activity.teacher_actions)
        if activity.student_actions:
            _sub_bullets(doc, "Students:", activity.student_actions)

    _bullet_section(doc, "Homework", plan.homework)
    _bullet_section(doc, "Assessment ideas", plan.assessment_ideas)
    _bullet_section(doc, "References", plan.references)

    if plan.outcome_codes_covered:
        foot = doc.add_paragraph()
        f_run = foot.add_run("Outcomes covered: " + ", ".join(plan.outcome_codes_covered))
        f_run.italic = True
        f_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _add_heading(doc, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _bullet_section(doc, heading: str, items: list[str]) -> None:
    if not items:
        return
    _add_heading(doc, heading, level=1)
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _sub_bullets(doc, label: str, items: list[str]) -> None:
    if not items:
        return
    label_par = doc.add_paragraph()
    l_run = label_par.add_run(label)
    l_run.bold = True
    for item in items:
        doc.add_paragraph(item, style="List Bullet 2")
