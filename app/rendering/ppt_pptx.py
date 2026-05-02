from io import BytesIO

from pptx import Presentation
from pptx.util import Inches, Pt

from app.llm.schemas.ppt import PPTOutlineOutput, Slide, SlideType


def render_ppt_outline(deck: PPTOutlineOutput) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    _add_title_slide(prs, deck)
    for slide in deck.slides[1:]:
        _add_content_slide(prs, slide)

    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def _add_title_slide(prs: Presentation, deck: PPTOutlineOutput) -> None:
    first = deck.slides[0]
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = first.title
    subtitle = slide.placeholders[1]
    subtitle.text = f"Class {deck.class_level}  ·  {deck.subject}  ·  {deck.chapter_title}"
    if first.speaker_notes:
        slide.notes_slide.notes_text_frame.text = first.speaker_notes


def _add_content_slide(prs: Presentation, slide_spec: Slide) -> None:
    layout = prs.slide_layouts[1]  # title + content
    slide = prs.slides.add_slide(layout)
    title_prefix = _title_prefix(slide_spec.slide_type)
    slide.shapes.title.text = f"{title_prefix}{slide_spec.title}"

    body = slide.placeholders[1].text_frame
    body.clear()
    for index, bullet in enumerate(slide_spec.bullets):
        paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        for run in paragraph.runs:
            run.font.size = Pt(22)

    if slide_spec.speaker_notes:
        slide.notes_slide.notes_text_frame.text = slide_spec.speaker_notes


def _title_prefix(slide_type: SlideType) -> str:
    mapping = {
        SlideType.OBJECTIVES: "Objectives — ",
        SlideType.EXAMPLE: "Example — ",
        SlideType.ACTIVITY: "Activity — ",
        SlideType.QUICK_CHECK: "Quick check — ",
        SlideType.SUMMARY: "Summary — ",
    }
    return mapping.get(slide_type, "")
